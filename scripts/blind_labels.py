"""Compare the blind labels (T112, part two) with the rules and the published severities (T113).

Reads the filled sheet once, copies the labels into datasets/labels/t112_blind_labels.csv so
the result can be reproduced from committed files, and writes results/blind_label_check.md.

    python scripts/blind_labels.py ../T112_blind_labels.xlsx
"""

from __future__ import annotations

import csv
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
LABELS = ROOT / "datasets" / "labels"
BANDS = ["high_or_critical", "medium", "non_security_or_low"]


def kappa(pairs):
    n = len(pairs)
    observed = sum(a == b for a, b in pairs) / n
    left, right = Counter(a for a, _ in pairs), Counter(b for _, b in pairs)
    expected = sum(left[b] * right[b] for b in BANDS) / (n * n)
    return round((observed - expected) / (1 - expected), 4) if expected < 1 else None


def import_sheet(path: Path) -> None:
    from openpyxl import load_workbook

    ws = load_workbook(path)["Label these"]
    with open(LABELS / "t112_blind_labels.csv", "w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(["item", "label", "confidence", "note"])
        for row in ws.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            if row[4] not in BANDS:
                raise SystemExit(f"item {row[0]} has no valid label")
            writer.writerow([row[0], row[4], row[5], row[6] or ""])


def main() -> int:
    from terraguard.feature_builder import load_rows
    from terraguard.model_trainer import labels_for

    if len(sys.argv) > 1:
        import_sheet(Path(sys.argv[1]))

    rows = load_rows(ROOT / "datasets/findings.csv", ROOT / "datasets/observations.csv")
    rule = {(r.run_id, r.finding_id): (b, r) for r, b in zip(rows, labels_for(rows), strict=True)}
    with open(ROOT / "datasets/external_severity_map.csv", encoding="utf-8") as fh:
        external = {
            (m["control_id"], m["normalised_event"], m["resource_type"]): m
            for m in csv.DictReader(fh)
        }
    with open(LABELS / "t112_sample.csv", encoding="utf-8") as fh:
        sample = {int(s["item"]): (s["run_id"], s["finding_id"]) for s in csv.DictReader(fh)}
    with open(LABELS / "t112_blind_labels.csv", encoding="utf-8") as fh:
        blind = {int(b["item"]): b for b in csv.DictReader(fh)}

    table, pairs, ext_pairs = [], [], []
    for item in sorted(sample):
        band, row = rule[sample[item]]
        key = (row.control_id or "-", row.normalised_event, row.resource_type)
        ref = external.get(key) or external.get((key[0], key[1], "*")) or {}
        mine = blind[item]["label"]
        pairs.append((band, mine))
        if ref.get("external_band"):
            ext_pairs.append((mine, ref["external_band"]))
        table.append(
            (
                item,
                key,
                band,
                mine,
                ref.get("external_band") or "-",
                blind[item]["confidence"],
                blind[item]["note"],
            )
        )

    agree = sum(a == b for a, b in pairs)
    disagreements = [t for t in table if t[2] != t[3]]
    lines = [
        "# Blind labels against the rules (T112 part two, T113)",
        "",
        "Twenty findings, one per finding type, labelled by the project author on 17 Oct without "
        "the rule's band, the rule number or the control id, using a plain impact scale rather "
        "than the rule ladder (`scripts/make_label_sheet.py`).",
        "",
        f"Agreement with the rules: {agree} of {len(pairs)} ({agree / len(pairs):.0%}), "
        f"Cohen's kappa {kappa(pairs)}.",
        "",
        f"Where a published severity exists ({len(ext_pairs)} items), the blind labels agree "
        f"with it on {sum(a == b for a, b in ext_pairs)}.",
        "",
        "| # | Finding type | Rule | Blind label | Published | Sure (1-3) |",
        "|---:|---|---|---|---|---:|",
    ]
    for item, key, band, mine, ext, sure, _ in table:
        mark = "" if band == mine else " (differs)"
        lines.append(f"| {item} | {key[0]} {key[2]} | {band} | {mine}{mark} | {ext} | {sure} |")
    lines += ["", "## Disagreements and decisions (T113)", ""]
    for item, key, band, mine, ext, _sure, note in disagreements:
        lines.append(
            f"- **#{item} {key[0]} on {key[2]}**: rule {band}, blind {mine}, "
            f"published {ext}. Labeller's reason: {note}"
        )
    lines += [
        "",
        "The rules are not changed in response. A disagreement is recorded as evidence about the",
        "rubric, not fixed after the fact, because tuning the rules to these twenty labels would",
        "make them the training set for the next comparison.",
        "",
        "Limitation: the labeller is the project author, so this is a weaker check than an",
        "independent reviewer. It covers finding types no published standard rates.",
        "",
    ]
    (ROOT / "results" / "blind_label_check.md").write_text("\n".join(lines), encoding="utf-8")
    print("\n".join(lines))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
