"""Finalise W9: regenerate results/, remove the one-off runners, update all six sheets.

The one-off runners existed only because the Linux workspace lost access to the files. Their
useful parts are now in scripts/build_results.py, which is a real pipeline step, and their
outputs live in results/ where they are committed. Deleting them is safe; this script checks
nothing references them first.
"""

from __future__ import annotations

import datetime as dt
import subprocess
import sys
from copy import copy
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT.parent / "TerraGuard_Build_Tracker.xlsx"
LOG = ROOT / "build/finalise_log.txt"
_lines: list[str] = []

ONE_OFFS = ["run_w8.py", "run_w8b.py", "run_w8c.py", "run_w9.py", "run_w9b.py", "run_w9c.py"]


def say(text: str = "") -> None:
    print(text)
    _lines.append(text)
    LOG.parent.mkdir(exist_ok=True)
    LOG.write_text("\n".join(_lines), encoding="utf-8")


def run(label, args, timeout=900):
    proc = subprocess.run(args, cwd=str(ROOT), capture_output=True, text=True, timeout=timeout)
    out = ((proc.stdout or "") + (proc.stderr or "")).strip()
    say(f"\n=== {label}  exit {proc.returncode}")
    say("\n".join(out.splitlines()[-15:]) or "(no output)")
    return proc.returncode


def main() -> int:
    from openpyxl import load_workbook

    say(f"date {dt.date.today().isoformat()}")

    # ---- 1. regenerate results/ from the committed dataset -------------------------
    if run("build_results", [sys.executable, str(ROOT / "scripts/build_results.py")]):
        say("!! results generation failed - nothing deleted, tracker untouched")
        return 1

    # ---- 2. check nothing references the one-offs before deleting them --------------
    say("\n=== checking for references to the one-off runners")
    searchable = list(ROOT.rglob("*.py")) + list(ROOT.rglob("*.md")) + list(ROOT.rglob("*.yml"))
    referenced = []
    for name in ONE_OFFS:
        stem = name[:-3]
        for f in searchable:
            # Skip the one-offs themselves, and this script - its own ONE_OFFS list names
            # all six, which would otherwise count as six references and block every run.
            if (
                f.name in ONE_OFFS
                or f.name == Path(__file__).name
                or "build" in f.parts
                or ".git" in f.parts
            ):
                continue
            try:
                if stem in f.read_text(encoding="utf-8", errors="ignore"):
                    referenced.append((name, str(f.relative_to(ROOT))))
            except OSError:
                continue
    if referenced:
        for name, where in referenced:
            say(f"  {name} still referenced by {where}")
        say("!! not deleting anything while references exist")
        return 1
    say("  no references found")

    removed = []
    for name in ONE_OFFS:
        target = ROOT / "scripts" / name
        if target.exists():
            target.unlink()
            removed.append(name)
    for log in (ROOT / "build").glob("*_log.txt"):
        if log.name != LOG.name:
            log.unlink()
    for stale in (
        "model_comparison.md",
        "model_features.csv",
        "training_report.json",
        "split_verification.json",
        "leakage_matrix.json",
        "data_quality.json",
        "duplicate_audit.json",
    ):
        old = ROOT / "build" / stale
        if old.exists():
            old.unlink()
    say(f"  removed {len(removed)} one-off runner(s): {removed}")
    say("  removed superseded copies from build/ - results/ is the committed home now")

    # ---- 3. gates -------------------------------------------------------------------
    paths = ["terraguard", "scripts", "tests", "scenarios"]
    run("black", [sys.executable, "-m", "black", "--line-length", "100", *paths, "-q"])
    lint_rc = run(
        "ruff", [sys.executable, "-m", "ruff", "check", "--output-format=concise", *paths[:3]]
    )
    fmt_rc = run(
        "black --check",
        [sys.executable, "-m", "black", "--check", "--line-length", "100", *paths],
    )
    test_rc = run("pytest", [sys.executable, "-m", "pytest", "-q"])
    if lint_rc or fmt_rc or test_rc:
        say(f"\n!! lint {lint_rc} format {fmt_rc} tests {test_rc} - tracker not touched")
        return 1

    # ---- 4. all six sheets ----------------------------------------------------------
    wb = load_workbook(BOOK)
    tt, cl, rr, ms, dash, hu = (
        wb["Task Tracker"],
        wb["Change Log"],
        wb["Risk Register"],
        wb["Milestones"],
        wb["Dashboard"],
        wb["How to use"],
    )
    idx = {
        str(tt.cell(r, 1).value).strip(): r
        for r in range(5, tt.max_row + 1)
        if tt.cell(r, 1).value and tt.cell(r, 7).value
    }

    # Task Tracker: every note that cited build/ pointed at an ignored directory.
    repointed = 0
    for r in range(5, tt.max_row + 1):
        note = tt.cell(r, 12).value
        if isinstance(note, str) and "build/" in note:
            tt.cell(r, 12).value = note.replace("build/", "results/")
            repointed += 1
    say(f"\n  repointed {repointed} task note(s) from build/ to results/")

    if "T115" in idx:
        tt.cell(idx["T115"], 12).value = (
            "datasets/reproducibility.json: seed, Python 3.14.3, sklearn 1.9.1, numpy 2.4.3, "
            "scipy 1.18.1, git commit, region. Terraform absent - no AWS run yet. "
            "scripts/build_results.py regenerates every results/ artefact from the committed "
            "dataset in one command."
        )

    # Risk Register: evidence that is gitignored is not evidence.
    used = {rr.cell(x, 1).value for x in range(4, rr.max_row + 1)}
    if "R21" not in used:
        r = rr.max_row + 1
        for c in range(1, 8):
            rr.cell(r, c)._style = copy(rr.cell(r - 1, c)._style)
        for c, v in enumerate(
            [
                "R21",
                "Result artefacts cited as evidence are not committed",
                "Medium",
                "Low",
                "Found during the W9 clean-up: every result was written to build/, which .gitignore "
                "excludes, while the tracker and methodology notes cited those paths. A supervisor "
                "cloning the repository would have found nothing. Results now live in results/, "
                "which is committed, and scripts/build_results.py regenerates them. The encoder "
                "pickle stays in build/, consistent with models/*.joblib already being ignored.",
                "T115",
                "Mitigated",
            ],
            1,
        ):
            rr.cell(r, c).value = v
        say(f"  R21 added at row {r}")

    for x in range(4, ms.max_row + 1):
        if ms.cell(x, 1).value == "M7":
            ms.cell(x, 2).value = (
                "LR + RF trained and compared on the offline set, results committed under "
                "results/; a rules-versus-models claim still needs external labels (T112)"
            )

    dash.cell(13, 8).value = (
        "P7 comparison complete, results committed; P8 reporting begun - generator, predictor, "
        "three writers and the redaction gate; D05, D08"
    )

    hu.cell(20, 1).value = "Where the evidence lives"
    hu.cell(20, 2).value = (
        "results/ holds every figure the tracker and the write-up cite, and it is committed. "
        "build/ is scratch and is gitignored - nothing cited as evidence belongs there. "
        "scripts/build_results.py regenerates the whole of results/ from the committed dataset "
        "in one command."
    )

    row = max(x for x in range(5, cl.max_row + 1) if cl.cell(x, 1).value) + 1
    for what, affected in [
        (
            "Every result was being written to build/, which .gitignore excludes, while the tracker "
            "and the methodology notes cited those exact paths as evidence. Anyone cloning the "
            "repository would have found the citations pointing at nothing. Results moved to "
            "results/, which is committed, and the notes repointed.",
            "T115",
        ),
        (
            "The six one-off scripts written to work around the broken workspace have been deleted "
            "and their useful parts consolidated into scripts/build_results.py, which regenerates "
            "every result artefact from the committed dataset in one command. Report generation "
            "was also removed from model_trainer, because two places writing a comparison file "
            "produced two files saying different things.",
            "T085, T115",
        ),
    ]:
        for c, v in enumerate([dt.date(2026, 9, 16), "W9", what, affected], 1):
            cell = cl.cell(row, c)
            cell._style = copy(cl.cell(row - 1, c)._style)
            cell.value = v
        cl.cell(row, 1).number_format = "dd mmm yyyy"
        row += 1

    wb.calculation.fullCalcOnLoad = True
    wb.save(BOOK)
    say(f"\nsaved {BOOK}")

    run("check_tracker", [sys.executable, str(ROOT / "scripts/check_tracker.py"), str(BOOK)])
    run("git status", ["git", "status", "--short"])
    say("\nDONE")
    return 0


if __name__ == "__main__":
    say("starting finalise")
    try:
        raise SystemExit(main())
    except SystemExit:
        raise
    except BaseException:
        import traceback

        say("!! failed")
        say(traceback.format_exc()[-2500:])
        raise SystemExit(1) from None
