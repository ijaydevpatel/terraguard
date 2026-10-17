"""Build the blind labelling sheet for T112 and the sample file that links it back to findings.

One finding per finding type: all eighteen security types and two non-security ones, drawn with
a fixed seed and shuffled. The sheet shows what was found and where, never the rule's band, the
rule number or the control id, and the guidance is a plain impact scale rather than the rule
ladder, which would only teach the labeller to reproduce the rules.

    python scripts/make_label_sheet.py ../T112_blind_labels.xlsx
"""

from __future__ import annotations

import csv
import random
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
SEED = 20261017
SAMPLE = ROOT / "datasets" / "labels" / "t112_sample.csv"
CHOICES = ["high_or_critical", "medium", "non_security_or_low"]

GUIDE = [
    (
        "high_or_critical",
        "Someone outside could read secrets or data, get in, or act with more "
        "power than intended, now or with one easy step. Or the record that would show an attack "
        "is gone.",
    ),
    (
        "medium",
        "Weakens a protection or makes a mistake more likely, but nothing is exposed and "
        "an attacker still needs something else first.",
    ),
    (
        "non_security_or_low",
        "Housekeeping. Labels, descriptions, names. Nothing about access, "
        "secrets or detection changes.",
    ),
]


def main() -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    from terraguard.feature_builder import load_rows

    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT.parent / "T112_blind_labels.xlsx"
    rows = load_rows(ROOT / "datasets/findings.csv", ROOT / "datasets/observations.csv")
    by_type = defaultdict(list)
    for row in rows:
        by_type[(row.control_id or "-", row.normalised_event, row.resource_type)].append(row)

    rng = random.Random(SEED)
    security = sorted(k for k in by_type if k[0] != "-")
    other = sorted(k for k in by_type if k[0] == "-")
    chosen = security + rng.sample(other, 20 - len(security))
    picks = [rng.choice(sorted(by_type[k], key=lambda r: r.finding_id)) for k in chosen]
    rng.shuffle(picks)

    SAMPLE.parent.mkdir(parents=True, exist_ok=True)
    with open(SAMPLE, "w", encoding="utf-8", newline="") as fh:
        writer = csv.writer(fh, lineterminator="\n")
        writer.writerow(["item", "run_id", "finding_id"])
        for i, row in enumerate(picks, 1):
            writer.writerow([i, row.run_id, row.finding_id])

    wb = Workbook()
    ws = wb.active
    ws.title = "Label these"
    head = [
        "#",
        "Where it was found",
        "Resource",
        "What changed or was found",
        "Your label",
        "How sure (1-3)",
        "Note (optional)",
    ]
    ws.append(head)
    for i, row in enumerate(picks, 1):
        where = {
            "repository": "Git repository",
            "workflow": "GitHub Actions workflow",
            "backend_config": "Terraform backend settings",
            "terraform_plan_json": "Live cloud vs Terraform (drift)",
        }.get(row.evidence_source, row.evidence_source)
        resource = row.resource_address or row.resource_type
        # Corpus paths carry the build machine's folders and the scenario's own name, which
        # would hint at the answer. Only the path inside the repository is shown.
        if "/.github/" in resource:
            resource = ".github/" + resource.split("/.github/", 1)[1]
        ws.append([i, where, resource, row.detail, "", "", ""])

    bold, fill = Font(bold=True), PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = bold
    for r in range(2, len(picks) + 2):
        for c in (5, 6, 7):
            ws.cell(r, c).fill = fill
        for c in range(1, 8):
            ws.cell(r, c).alignment = Alignment(wrap_text=True, vertical="top")
    for col, width in zip("ABCDEFG", (5, 26, 34, 60, 22, 14, 30), strict=True):
        ws.column_dimensions[col].width = width
    labels = DataValidation(type="list", formula1='"' + ",".join(CHOICES) + '"', allow_blank=True)
    sure = DataValidation(type="list", formula1='"1,2,3"', allow_blank=True)
    ws.add_data_validation(labels)
    ws.add_data_validation(sure)
    labels.add(f"E2:E{len(picks) + 1}")
    sure.add(f"F2:F{len(picks) + 1}")
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("How to label")
    guide.append(["How to label"])
    guide["A1"].font = Font(bold=True, size=13)
    for line in [
        "Read each row and pick the label that fits the harm it could cause, using only the "
        "scale below. Go with your own judgement; there is no answer key on this sheet.",
        "Do not look up the project's severity rules, the tracker or the code while labelling. "
        "The point is a judgement that did not come from the rules.",
        "How sure: 1 = guessing, 2 = fairly sure, 3 = certain. Use the note for anything odd.",
        "Takes about 15 minutes. Save the file when done.",
    ]:
        guide.append([line])
    guide.append([])
    guide.append(["Label", "Use it when"])
    for cell in guide[guide.max_row]:
        cell.font = bold
    for name, text in GUIDE:
        guide.append([name, text])
    guide.column_dimensions["A"].width = 22
    guide.column_dimensions["B"].width = 90
    for row in guide.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb.save(out)
    print(f"{len(picks)} findings written to {out}; sample ids in {SAMPLE}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
